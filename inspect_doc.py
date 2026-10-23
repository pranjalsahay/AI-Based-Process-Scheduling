import docx

doc = docx.Document('WattWise_BCSE302P_Mini_Project_Report.docx')
with open('extracted_structure.txt', 'w', encoding='utf-8') as out:
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            out.write(f"P[{i}] (Style: {p.style.name}): {txt}\n")
    for t_idx, tbl in enumerate(doc.tables):
        out.write(f"\n--- Table {t_idx} ({len(tbl.rows)}x{len(tbl.columns)}) ---\n")
        for r_idx, row in enumerate(tbl.rows[:3]):
            cells = [c.text.replace('\n', ' ') for c in row.cells]
            out.write(f"  Row {r_idx}: {' | '.join(cells)}\n")

print("Wrote extracted_structure.txt successfully!")
